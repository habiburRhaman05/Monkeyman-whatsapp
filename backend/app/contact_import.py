"""Parse CSV / XLSX uploads into cleaned phone-number rows."""

import csv
import io
import json
import re
from typing import Any

_PHONE_RE = re.compile(r"\d{6,15}$")

_FIRST_NAME_HINTS = {"first_name", "firstname", "first", "given_name", "givenname", "prenom", "first_names"}
_LAST_NAME_HINTS = {"last_name", "lastname", "last", "surname", "family_name", "familyname", "last_names"}
_NAME_HINTS = {"name", "full_name", "fullname", "contact", "contact_name", "nombre", "display_name", "displayname", "customer_name", "customer"}
_EMAIL_HINTS = {"email", "e_mail", "email_address", "emailaddress", "mail", "correo", "e-mail", "emails", "email_id", "emailid", "contact_email"}
_PHONE_HINTS = {"phone", "phone_number", "phonenumber", "mobile", "cell", "telephone", "tel", "number", "whatsapp", "phone_no", "mobile_number", "contact_number", "cell_phone", "cellphone"}
_CC_HINTS = {"country_code", "countrycode", "cc", "dial_code", "prefix", "phone_code"}
_COUNTRY_NAME_HINTS = {"country", "country_name", "countryname", "nation", "location_country"}
_WA_COPY_HINTS = {"whatsapp_copy", "whatsappcopy", "wa_copy", "wacopy", "whatsapp_message", "wa_message", "message_copy", "messagecopy", "copy", "message_text", "sms_copy"}

COUNTRY_TO_CODE: dict[str, str] = {
    "us": "1", "usa": "1", "united states": "1", "united states of america": "1", "america": "1",
    "ca": "1", "canada": "1",
    "uk": "44", "gb": "44", "united kingdom": "44", "england": "44", "britain": "44", "great britain": "44",
    "bd": "880", "bangladesh": "880",
    "in": "91", "india": "91",
    "pk": "92", "pakistan": "92",
    "de": "49", "germany": "49", "deutschland": "49",
    "fr": "33", "france": "33",
    "es": "34", "spain": "34",
    "it": "39", "italy": "39", "italia": "39",
    "nl": "31", "netherlands": "31", "holland": "31",
    "be": "32", "belgium": "32",
    "pt": "351", "portugal": "351",
    "se": "46", "sweden": "46",
    "no": "47", "norway": "47",
    "dk": "45", "denmark": "45",
    "fi": "358", "finland": "358",
    "pl": "48", "poland": "48",
    "at": "43", "austria": "43",
    "ch": "41", "switzerland": "41",
    "ie": "353", "ireland": "353",
    "au": "61", "australia": "61",
    "nz": "64", "new zealand": "64",
    "br": "55", "brazil": "55",
    "mx": "52", "mexico": "52",
    "ae": "971", "uae": "971", "united arab emirates": "971",
    "sa": "966", "saudi arabia": "966",
    "sg": "65", "singapore": "65",
    "my": "60", "malaysia": "60",
    "ph": "63", "philippines": "63",
    "ng": "234", "nigeria": "234",
    "za": "27", "south africa": "27",
    "ke": "254", "kenya": "254",
    "gh": "233", "ghana": "233",
    "eg": "20", "egypt": "20",
    "tr": "90", "turkey": "90", "türkiye": "90",
    "jp": "81", "japan": "81",
    "kr": "82", "south korea": "82", "korea": "82",
    "cn": "86", "china": "86",
    "ru": "7", "russia": "7",
    "id": "62", "indonesia": "62",
    "th": "66", "thailand": "66",
    "vn": "84", "vietnam": "84",
    "lk": "94", "sri lanka": "94",
    "np": "977", "nepal": "977",
    "mm": "95", "myanmar": "95",
    "co": "57", "colombia": "57",
    "ar": "54", "argentina": "54",
    "cl": "56", "chile": "56",
    "pe": "51", "peru": "51",
}


def country_name_to_code(name: str | None) -> str | None:
    if not name:
        return None
    return COUNTRY_TO_CODE.get(name.strip().lower())


def guess_mapping(columns: list[str]) -> dict[str, str | None]:
    lower = {c: c.strip().lower().replace(" ", "_") for c in columns}
    mapping: dict[str, str | None] = {
        "first_name": None,
        "last_name": None,
        "name": None,
        "email": None,
        "phone": None,
        "country_code": None,
        "country_name": None,
        "whatsapp_copy": None,
    }
    for orig, norm in lower.items():
        if norm in _FIRST_NAME_HINTS and mapping["first_name"] is None:
            mapping["first_name"] = orig
        elif norm in _LAST_NAME_HINTS and mapping["last_name"] is None:
            mapping["last_name"] = orig
        elif norm in _NAME_HINTS and mapping["name"] is None:
            mapping["name"] = orig
        elif norm in _EMAIL_HINTS and mapping["email"] is None:
            mapping["email"] = orig
        elif norm in _PHONE_HINTS and mapping["phone"] is None:
            mapping["phone"] = orig
        elif norm in _CC_HINTS and mapping["country_code"] is None:
            mapping["country_code"] = orig
        elif norm in _COUNTRY_NAME_HINTS and mapping["country_name"] is None:
            mapping["country_name"] = orig
        elif norm in _WA_COPY_HINTS and mapping["whatsapp_copy"] is None:
            mapping["whatsapp_copy"] = orig
    return mapping


def _read_csv(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("latin-1")
    reader = csv.DictReader(io.StringIO(text))
    columns = list(reader.fieldnames or [])
    rows = [dict(r) for r in reader]
    return columns, rows


def _read_xlsx(data: bytes) -> tuple[list[str], list[dict[str, str]]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    if ws is None:
        return [], []
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not all_rows:
        return [], []
    columns = [str(c).strip() if c is not None else f"col_{i}" for i, c in enumerate(all_rows[0])]
    rows: list[dict[str, str]] = []
    for row in all_rows[1:]:
        d: dict[str, str] = {}
        for i, val in enumerate(row):
            if i < len(columns) and val is not None:
                d[columns[i]] = str(val).strip()
        if any(d.values()):
            rows.append(d)
    return columns, rows


def parse_file(data: bytes, filename: str) -> tuple[list[str], list[dict[str, str]]]:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext in ("xlsx", "xls"):
        return _read_xlsx(data)
    return _read_csv(data)


def clean_phone(raw: str | None, row_cc: str | None, default_cc: str | None) -> str | None:
    """Clean and normalize a phone number.

    Rules:
    - Strip all non-digits.
    - If the number already starts with a country code (length > 10), keep it as-is.
    - Only prepend a country code when one is explicitly provided (per-row or default)
      AND the number is short enough to be a local number (<=10 digits).
    - Never silently add a prefix if no country code is given.
    """
    if not raw:
        return None
    s = str(raw)
    # Excel scientific notation: 8.8E+11 → 880000000000
    if "e" in s.lower() or "E" in s:
        try:
            s = str(int(float(s)))
        except (ValueError, OverflowError):
            pass
    # Keep the leading + info before stripping
    had_plus = s.strip().startswith("+")
    s = re.sub(r"[^\d]", "", s)
    if s.startswith("00"):
        s = s[2:]

    cc = re.sub(r"[^\d]", "", row_cc or "") if row_cc else None
    if not cc:
        cc = re.sub(r"[^\d]", "", default_cc or "") if default_cc else None

    # Only prepend country code if:
    # 1. A country code was explicitly provided (row-level or default)
    # 2. The number doesn't already start with that code
    # 3. The number looks like a local number (<=10 digits)
    # 4. The number didn't start with "+" (which means it already has international format)
    if cc and not had_plus and not s.startswith(cc) and len(s) <= 10:
        s = cc + s

    if not _PHONE_RE.match(s):
        return None
    return s


def process_rows(
    rows: list[dict[str, str]],
    mapping: dict[str, str | None],
    default_country_code: str | None,
) -> tuple[list[dict[str, Any]], int, int]:
    """Return (cleaned_rows, duplicate_count, invalid_count)."""
    first_name_col = mapping.get("first_name")
    last_name_col = mapping.get("last_name")
    name_col = mapping.get("name")
    email_col = mapping.get("email")
    phone_col = mapping.get("phone")
    cc_col = mapping.get("country_code")
    country_name_col = mapping.get("country_name")
    mapped_keys = {v for v in mapping.values() if v}

    seen: set[str] = set()
    cleaned: list[dict[str, Any]] = []
    dups = 0
    invalid = 0

    for row in rows:
        raw_phone = row.get(phone_col) if phone_col else None

        # Resolve country code: explicit cc column > country name column > default
        row_cc = row.get(cc_col) if cc_col else None
        if not row_cc and country_name_col:
            country_val = row.get(country_name_col)
            row_cc = country_name_to_code(country_val)

        phone = clean_phone(raw_phone, row_cc, default_country_code)
        if not phone:
            invalid += 1
            continue
        if phone in seen:
            dups += 1
            continue
        seen.add(phone)

        first_name = (row.get(first_name_col) or "").strip() if first_name_col else None
        last_name = (row.get(last_name_col) or "").strip() if last_name_col else None
        name = (row.get(name_col) or "").strip() if name_col else None
        email = (row.get(email_col) or "").strip() if email_col else None

        # Auto-derive name from first + last if no explicit name column
        if not name and (first_name or last_name):
            name = " ".join(filter(None, [first_name, last_name]))

        # Auto-derive first_name from full name if no first_name column
        if not first_name and name:
            parts = name.split()
            first_name = parts[0] if parts else None
            if not last_name and len(parts) > 1:
                last_name = " ".join(parts[1:])

        extra_cols = {k: v for k, v in row.items() if k not in mapped_keys and v}

        final_cc = (row_cc or default_country_code or "").strip() or None

        cleaned.append({
            "first_name": first_name or None,
            "last_name": last_name or None,
            "name": name or None,
            "email": email or None,
            "country_code": final_cc,
            "phone": phone,
            "extra": json.dumps(extra_cols) if extra_cols else None,
        })

    return cleaned, dups, invalid


# ── Template variable rendering ─────────────────────

_VAR_RE = re.compile(r"\{\{(\w+)\}\}")


def render_template(template: str, contact: dict) -> str:
    """Replace ``{{var}}`` placeholders with contact field values.

    Built-in fields (``first_name``, ``last_name``, ``name``, ``email``, ``phone``)
    are taken from the contact record directly.
    Everything else is looked up in the ``extra`` JSON object.
    A variable with no matching value is replaced with an empty string.
    """
    extra: dict = {}
    raw_extra = contact.get("extra")
    if raw_extra:
        try:
            extra = json.loads(raw_extra) if isinstance(raw_extra, str) else raw_extra
        except (json.JSONDecodeError, TypeError):
            pass

    DIRECT_FIELDS = ("first_name", "last_name", "name", "email", "phone", "country_code")

    def _replace(m: re.Match) -> str:
        key = m.group(1)
        if key in DIRECT_FIELDS:
            return str(contact.get(key) or "")
        return str(extra.get(key, ""))

    return _VAR_RE.sub(_replace, template)
